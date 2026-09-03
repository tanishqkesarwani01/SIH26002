"""
Script to generate SIH 2026 Official Idea PPTX and PDF Presentations
for Problem Statement: AI & GIS-Driven Logistics Resilience Platform for Northeast India (NER-LRP).
Follows the official 6-slide SIH template guidelines strictly.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.pdfgen import canvas


# ==============================================================================
# 1. GENERATE POWERPOINT PRESENTATION (.PPTX)
# ==============================================================================
def create_pptx(filename="SIH2026_NER_Logistics_Resilience_Platform.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY = RGBColor(15, 23, 42)      # #0f172a
    C_BLUE = RGBColor(30, 58, 138)     # #1e3a8a
    C_TEAL = RGBColor(13, 148, 136)    # #0d9488
    C_EMERALD = RGBColor(16, 185, 129) # #10b981
    C_ORANGE = RGBColor(245, 158, 11)  # #f59e0b
    C_LIGHT_BG = RGBColor(248, 250, 252) # #f8fafc
    C_CARD_BG = RGBColor(255, 255, 255)
    C_BORDER = RGBColor(203, 213, 225)
    C_TEXT_DARK = RGBColor(30, 41, 59)
    C_TEXT_MUTED = RGBColor(100, 116, 139)
    C_WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title_text, category_badge="SIH 2026 Idea Submission"):
        # Header background
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = C_NAVY
        header_box.line.fill.background()

        # Accent line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
        line.fill.solid()
        line.fill.fore_color.rgb = C_TEAL
        line.line.fill.background()

        # Title Text
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.8)
        tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = title_text.upper()
        p.font.bold = True
        p.font.size = Pt(22)
        p.font.color.rgb = C_WHITE
        p.font.name = 'Arial'

        # Right badge
        p2 = tf.add_paragraph()
        p2.text = f"SMART INDIA HACKATHON 2026  |  {category_badge}"
        p2.font.size = Pt(10)
        p2.font.color.rgb = C_EMERALD
        p2.font.name = 'Arial'

        # Footer
        footer_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.1), Inches(13.333), Inches(0.4))
        footer_box.fill.solid()
        footer_box.fill.fore_color.rgb = C_NAVY
        footer_box.line.fill.background()
        ftf = footer_box.text_frame
        ftf.margin_left = Inches(0.8)
        ftf.margin_top = Inches(0.08)
        fp = ftf.paragraphs[0]
        fp.text = "NER-LRP: AI & GIS-Driven Regional Logistics Control Tower for Northeast India  |  @SIH Idea Submission"
        fp.font.size = Pt(9)
        fp.font.color.rgb = C_TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 1: TITLE PAGE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Background
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY
    bg1.line.fill.background()

    # Title Card
    tbox = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.8))
    tbox.fill.solid()
    tbox.fill.fore_color.rgb = C_CARD_BG
    tbox.line.color.rgb = C_TEAL
    tbox.line.width = Pt(2)

    ttf = tbox.text_frame
    ttf.word_wrap = True
    ttf.margin_left = Inches(0.6)
    ttf.margin_right = Inches(0.6)
    ttf.margin_top = Inches(0.4)

    p = ttf.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEAL
    p.alignment = PP_ALIGN.CENTER

    p = ttf.add_paragraph()
    p.text = "NORTHEAST INDIA LOGISTICS RESILIENCE PLATFORM (NER-LRP)"
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = C_NAVY
    p.alignment = PP_ALIGN.CENTER

    p = ttf.add_paragraph()
    p.text = "AI & GIS-Driven Predictive Disruption Management & Resilient Logistics Control Tower"
    p.font.size = Pt(13)
    p.font.color.rgb = C_TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER

    # Divider line
    p = ttf.add_paragraph()
    p.text = "―" * 55
    p.font.size = Pt(10)
    p.font.color.rgb = C_BORDER
    p.alignment = PP_ALIGN.CENTER

    # Metadata Grid
    meta_items = [
        ("• Problem Statement ID:", "SIH26002"),
        ("• Problem Statement Title:", "Predictive Logistics Decision Support & Infrastructure Disruption Management in NER"),
        ("• Theme:", "Transportation & Logistics / Disaster Management / Smart Automation"),
        ("• PS Category:", "Software (AI/ML + Full-Stack Web + Offline-First PWA + GIS)"),
        ("• Team ID:", "[Your Team ID]"),
        ("• Team Name:", "[Your Team Name - Registered on SIH Portal]")
    ]

    for label, val in meta_items:
        p = ttf.add_paragraph()
        run1 = p.add_run()
        run1.text = f"{label} "
        run1.font.bold = True
        run1.font.size = Pt(12)
        run1.font.color.rgb = C_BLUE

        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(12)
        run2.font.color.rgb = C_TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 2: PROPOSED SOLUTION
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Proposed Solution: NER-LRP Control Tower")

    # Left Column: Detailed Explanation (Card)
    card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_CARD_BG
    card1.line.color.rgb = C_BORDER
    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_top = Inches(0.3)

    p = tf1.paragraphs[0]
    p.text = "Detailed Explanation of Proposed Solution"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_BLUE

    points1 = [
        ("Sense-Predict-Decide-Act Framework:", " Unified regional control tower integrating live satellite weather, terrain slope gradients, and road topologies."),
        ("Real Data Ingestion:", " Fetches live precipitation & 48h rain from Open-Meteo/IMD APIs across 15+ NER stations; integrates NASA SRTM 30m DEM slope angles."),
        ("Explainable AI Hazard Engine:", " Gradient Boosting model predicts 0-100 hazard risk with feature importance breakdown (Rain vs Slope vs Soil vs GSI History)."),
        ("Dynamic Corridor Rerouting:", " Automatically re-plans routes around blocked corridors (e.g., NH-27 bypass around NH-6 Sonapur) with live delta ETAs.")
    ]
    for h, desc in points1:
        p = tf1.add_paragraph()
        r1 = p.add_run()
        r1.text = f"• {h}"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_NAVY
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = C_TEXT_DARK

    # Right Column: Problem Fit & Innovation (Card)
    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    card2.fill.solid()
    card2.fill.fore_color.rgb = C_CARD_BG
    card2.line.color.rgb = C_BORDER
    tf2 = card2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = "How It Addresses Problem & Key Innovations"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_TEAL

    points2 = [
        ("Cargo-Criticality Weighting (Novel):", " 85% safety priority for Emergency Medicine & Oxygen vs 25% for Commercial Bulk. Eliminates life-or-death stranding."),
        ("Offline-First Field Officer PWA:", " Enables patrol officers in remote 0-network valleys to submit geo-photo incident reports cached locally in IndexedDB with auto-sync."),
        ("Multilingual Emergency Broadcaster:", " Translates road alerts into 6 regional languages (Assamese, Mizo, Manipuri, Bengali, Hindi, English) with synthesized voice audio."),
        ("1-Click Disruption Simulator:", " Allows live evaluators to trigger instant landslides (NH-6, NH-10, NH-29) to verify automated rerouting in real time.")
    ]
    for h, desc in points2:
        p = tf2.add_paragraph()
        r1 = p.add_run()
        r1.text = f"• {h}"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_TEAL
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = C_TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 3: TECHNICAL APPROACH
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Technical Approach & Architecture")

    # Top: Tech Stack Grid (3 Cards)
    tech_boxes = [
        ("Backend & Data Ingestion", ["• Python FastAPI (High-performance Async)", "• Open-Meteo & IMD Live Weather APIs", "• SRTM DEM Elevation & Slope Models", "• SQLite ACID Database & Spatial Schema"], C_BLUE, 0.8),
        ("AI / ML & Routing Engine", ["• Scikit-Learn Gradient Boosting Regressor", "• Multi-Factor Hazard Scoring (0-100)", "• Dijkstra Multi-Objective Graph Router", "• Cargo-Criticality Priority Cost Functions"], C_TEAL, 4.8),
        ("Frontend, GIS & PWA Suite", ["• React 18 + Vite + Tailwind CSS Dark UI", "• Leaflet GIS + React-Leaflet Map Layers", "• Offline PWA + IndexedDB Sync Queue", "• Web Speech API (6-Language Voice Audio)"], C_EMERALD, 8.8)
    ]

    for title, items, color, left_pos in tech_boxes:
        c = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(1.4), Inches(3.7), Inches(2.6))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = color
        c.line.width = Pt(1.5)
        ctf = c.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.2)
        ctf.margin_top = Inches(0.2)
        p = ctf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(12.5)
        p.font.color.rgb = color
        for item in items:
            p = ctf.add_paragraph()
            p.text = item
            p.font.size = Pt(9.5)
            p.font.color.rgb = C_TEXT_DARK

    # Bottom: Implementation Architecture Workflow Box
    arch_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.6))
    arch_box.fill.solid()
    arch_box.fill.fore_color.rgb = C_CARD_BG
    arch_box.line.color.rgb = C_BORDER
    atf = arch_box.text_frame
    atf.word_wrap = True
    atf.margin_left = Inches(0.3)
    atf.margin_top = Inches(0.2)

    p = atf.paragraphs[0]
    p.text = "End-to-End Implementation Methodology & Process Workflow"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_NAVY

    workflow_steps = [
        ("1. Real Ingestion:", " Queries live rain, wind & soil moisture from Open-Meteo for 15+ NER stations; pairs with SRTM 30m slope gradients."),
        ("2. ML Hazard Scoring:", " Gradient Boosting regressor calculates 0-100 composite risk score with explainable feature importance attribution."),
        ("3. Cargo Prioritization:", " Multi-objective Dijkstra router optimizes path based on cargo type (Medical: 85% safety; Bulk: 75% speed)."),
        ("4. Control Tower & PWA:", " Live GIS dark map renders color-coded corridors; offline PWA captures geo-photos in remote mountain valleys."),
        ("5. Voice & SMS Dispatch:", " Generates synthesized voice audio and WhatsApp/SMS alerts translated into 6 Northeast regional languages.")
    ]
    for step, desc in workflow_steps:
        p = atf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{step} "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_BLUE
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(10)
        r2.font.color.rgb = C_TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Feasibility, Viability & Risk Mitigation")

    f_cards = [
        ("Feasibility & Prototyping", [
            ("• Open Data Architecture:", " Fully built on free, open APIs (Open-Meteo, OpenStreetMap, NASA SRTM DEM). Zero paid API dependencies."),
            ("• Working End-to-End Prototype:", " Complete full-stack prototype active on local/cloud ports with 21 passing automated unit tests."),
            ("• Corridor-Scale Validation:", " Validated across major Northeast lifelines: NH-10 (Sikkim), NH-6 (Meghalaya-Barak), NH-29 (Nagaland), NH-27 (Assam Bypass).")
        ], C_EMERALD, 0.8),
        ("Potential Challenges & Risks", [
            ("• Government Fleet Data Locked:", " Live vehicle GPS and ULIP feeds require formal NDAs and ministry clearances."),
            ("• Cellular Blackouts in Remote Ghats:", " Mountain passes frequently suffer from zero 4G/5G mobile connectivity."),
            ("• Non-linear Himalayan Terrain:", " Flash floods and cloudbursts can cause sudden unforeseen landslides.")
        ], C_ORANGE, 4.8),
        ("Mitigation & Scalability Strategies", [
            ("• Modular Ingestion Adapter:", " Designed to plug into live ULIP / state transport APIs seamlessly when authorized."),
            ("• Offline-First PWA Local Storage:", " Stores field reports in IndexedDB; auto-syncs with idempotent protocol upon network recovery."),
            ("• Hybrid Physics + ML Scoring:", " Combines empirical Indian Road Congress (IRC) slope rules with GBDT predictive scoring.")
        ], C_BLUE, 8.8)
    ]

    for title, items, color, left_pos in f_cards:
        c = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(1.4), Inches(3.7), Inches(5.4))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = color
        c.line.width = Pt(1.5)
        ctf = c.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.3)

        p = ctf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = color

        for h, desc in items:
            p = ctf.add_paragraph()
            r1 = p.add_run()
            r1.text = f"{h} "
            r1.font.bold = True
            r1.font.size = Pt(10)
            r1.font.color.rgb = C_NAVY
            r2 = p.add_run()
            r2.text = desc
            r2.font.size = Pt(10)
            r2.font.color.rgb = C_TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 5: IMPACT AND BENEFITS
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Impact, Benefits & Stakeholder Value")

    # 4 Impact Quadrants
    quadrants = [
        ("👥 Target Audience Impact", [
            ("• MDoNER & SDMAs:", " Unified regional control tower across all 8 Northeast states for coordinated disaster response."),
            ("• Hospitals & Clinics:", " Guarantees zero stock-out of emergency oxygen, anti-venoms, and cold-chain vaccines."),
            ("• Transporters & Drivers:", " Hands-free multilingual audio alerts prevent getting trapped in 48-hour mountain jams.")
        ], C_BLUE, 0.8, 1.4),
        ("💰 Economic Benefits", [
            ("• Inflation Suppression:", " Prevents local price spikes (>20%) in isolated districts (Mizoram, Sikkim, Nagaland)."),
            ("• Fuel & Delay Savings:", " Cuts vehicle idling and turnaround times by 15-25% through planned bypass diversions."),
            ("• Fleet Asset Protection:", " Prevents multi-crore truck damages from rockfalls and sudden landslide collapses.")
        ], C_EMERALD, 6.8, 1.4),
        ("🌿 Social & Environmental Benefits", [
            ("• Supply Reliability:", " Uninterrupted food grain, PDS rations, and baby food delivery to isolated tribal hamlets."),
            ("• Reduced Carbon Burn:", " Eliminates thousands of hours of diesel engine idling during highway blockades."),
            ("• Rural Empowerment:", " Provides local field scouts and village heads a direct voice to central command via PWA.")
        ], C_TEAL, 0.8, 4.3),
        ("🛡️ System Resilience & PM GatiShakti", [
            ("• Emergency Corridor Identification:", " Uncovers secondary state highways and defense bypasses during primary closures."),
            ("• GatiShakti Alignment:", " Integrates multi-modal GIS layers to assist infrastructure planners in fortifying slope cuts."),
            ("• Scale beyond Hackathon:", " Scalable to all Himalayan mountain regions (Uttarakhand, Himachal, Jammu & Kashmir).")
        ], C_NAVY, 6.8, 4.3)
    ]

    for title, items, color, left_pos, top_pos in quadrants:
        c = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(top_pos), Inches(5.7), Inches(2.6))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = color
        c.line.width = Pt(1.5)
        ctf = c.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.2)

        p = ctf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(12.5)
        p.font.color.rgb = color

        for h, desc in items:
            p = ctf.add_paragraph()
            r1 = p.add_run()
            r1.text = f"{h} "
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = C_NAVY
            r2 = p.add_run()
            r2.text = desc
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = C_TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 6: RESEARCH AND REFERENCES
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Research, Citations & References")

    # Two Column Reference Cards
    ref_card1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    ref_card1.fill.solid()
    ref_card1.fill.fore_color.rgb = C_CARD_BG
    ref_card1.line.color.rgb = C_BORDER
    rtf1 = ref_card1.text_frame
    rtf1.word_wrap = True
    rtf1.margin_left = Inches(0.3)
    rtf1.margin_top = Inches(0.3)

    p = rtf1.paragraphs[0]
    p.text = "Geospatial Data & Scientific Literature"
    p.font.bold = True
    p.font.size = Pt(13.5)
    p.font.color.rgb = C_BLUE

    refs1 = [
        ("1. Geological Survey of India (GSI):", " National Landslide Susceptibility Mapping (NLSM) and Landslide Hazard Zonation (LHZ) protocols for Northeast India."),
        ("2. Academic ML Research (Eastern Himalayas):", " XGBoost / LightGBM models for landslide susceptibility in NE India (Dibang Valley & Nagaland, achieving AUC ~0.96 and ~0.89)."),
        ("3. NASA SRTM 30m Digital Elevation Model:", " High-resolution digital elevation data used for terrain slope and drainage gradient modeling across Himalayan highways."),
        ("4. Open-Meteo & IMD Meteorological Data:", " Public API endpoints providing live precipitation, 48-hour cumulative rainfall, and Doppler weather radar data.")
    ]
    for h, desc in refs1:
        p = rtf1.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{h}\n"
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_NAVY
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_DARK

    ref_card2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    ref_card2.fill.solid()
    ref_card2.fill.fore_color.rgb = C_CARD_BG
    ref_card2.line.color.rgb = C_BORDER
    rtf2 = ref_card2.text_frame
    rtf2.word_wrap = True
    rtf2.margin_left = Inches(0.3)
    rtf2.margin_top = Inches(0.3)

    p = rtf2.paragraphs[0]
    p.text = "Government Policies, Frameworks & Standards"
    p.font.bold = True
    p.font.size = Pt(13.5)
    p.font.color.rgb = C_TEAL

    refs2 = [
        ("5. PM GatiShakti National Master Plan:", " Ministry of Commerce & Industry digital planning mandate for multimodal connectivity and logistics infrastructure resilience."),
        ("6. High-Level Task Force (HLTF) NER Report:", " Official MDoNER reports detailing transport bottlenecks, logistics costs, and the 20% price inflation in remote hill states."),
        ("7. ULIP (Unified Logistics Interface Platform):", " National Logistics Policy framework integrating multi-ministry vehicle telematics and digital transport documentation."),
        ("8. OpenStreetMap (OSM) & OSRM Routing Engine:", " Open-source road network graph and multi-criteria routing algorithms customized for heavy truck transport.")
    ]
    for h, desc in refs2:
        p = rtf2.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{h}\n"
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_TEAL
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_DARK

    prs.save(filename)
    print(f"[SUCCESS] PPTX presentation successfully generated at: {filename}")


# ==============================================================================
# 2. GENERATE MATCHING LANDSCAPE PDF PRESENTATION (.PDF)
# ==============================================================================
def create_pdf_presentation(filename="SIH2026_NER_Logistics_Resilience_Platform_Presentation.pdf"):
    # 16:9 ratio landscape letter (792 x 445.5 pt)
    page_w, page_h = 792, 445.5
    doc = SimpleDocTemplate(
        filename,
        pagesize=(page_w, page_h),
        leftMargin=36,
        rightMargin=36,
        topMargin=28,
        bottomMargin=28
    )

    styles = getSampleStyleSheet()

    c_navy = colors.HexColor("#0f172a")
    c_blue = colors.HexColor("#1e3a8a")
    c_teal = colors.HexColor("#0d9488")
    c_dark = colors.HexColor("#1e293b")
    c_gray = colors.HexColor("#475569")
    c_light = colors.HexColor("#f8fafc")
    c_border = colors.HexColor("#cbd5e1")

    title_slide_title = ParagraphStyle(
        'TTitle', parent=styles['Heading1'],
        fontName='Helvetica-Bold', fontSize=18, leading=22,
        textColor=c_navy, alignment=1, spaceAfter=4
    )
    title_slide_sub = ParagraphStyle(
        'TSub', parent=styles['Normal'],
        fontName='Helvetica', fontSize=11, leading=14,
        textColor=c_teal, alignment=1, spaceAfter=8
    )

    slide_header_style = ParagraphStyle(
        'SHeader', parent=styles['Heading2'],
        fontName='Helvetica-Bold', fontSize=13, leading=16,
        textColor=colors.white
    )

    card_h_style = ParagraphStyle(
        'CardH', parent=styles['Heading3'],
        fontName='Helvetica-Bold', fontSize=10.5, leading=13,
        textColor=c_blue, spaceAfter=4
    )
    card_h_teal = ParagraphStyle(
        'CardHTeal', parent=card_h_style,
        textColor=c_teal
    )

    body_style = ParagraphStyle(
        'BStyle', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8, leading=10.5,
        textColor=c_dark
    )

    meta_lbl = ParagraphStyle(
        'MetaLbl', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=9.5, leading=13,
        textColor=c_blue
    )
    meta_val = ParagraphStyle(
        'MetaVal', parent=styles['Normal'],
        fontName='Helvetica', fontSize=9.5, leading=13,
        textColor=c_dark
    )

    def make_slide_header(title_text, slide_num):
        h_data = [
            [
                Paragraph(f"<b>{title_text.upper()}</b>", slide_header_style),
                Paragraph(f"<font size=7.5 color='#94a3b8'>SMART INDIA HACKATHON 2026 | Slide {slide_num} of 6</font>", ParagraphStyle('R', parent=slide_header_style, alignment=2))
            ]
        ]
        t = Table(h_data, colWidths=[520, 200])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), c_navy),
            ('PADDING', (0,0), (-1,-1), 6),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LINEBELOW', (0,-1), (-1,-1), 2, c_teal)
        ]))
        return t

    story = []

    # -------------------------------------------------------------
    # SLIDE 1: TITLE PAGE
    # -------------------------------------------------------------
    t_box_data = [
        [Paragraph("<b>SMART INDIA HACKATHON 2026</b>", ParagraphStyle('H', parent=title_slide_sub, fontSize=11, fontName='Helvetica-Bold'))],
        [Paragraph("<b>NORTHEAST INDIA LOGISTICS RESILIENCE PLATFORM (NER-LRP)</b>", title_slide_title)],
        [Paragraph("<b>AI & GIS-Driven Predictive Disruption Management & Resilient Logistics Control Tower</b>", title_slide_sub)],
        [
            Table([
                [Paragraph("<b>• Problem Statement ID:</b>", meta_lbl), Paragraph("SIH26002", meta_val)],
                [Paragraph("<b>• Problem Statement Title:</b>", meta_lbl), Paragraph("Predictive Logistics Decision Support & Disruption Management for NER", meta_val)],
                [Paragraph("<b>• Theme:</b>", meta_lbl), Paragraph("Transportation & Logistics / Disaster Management / Smart Automation", meta_val)],
                [Paragraph("<b>• PS Category:</b>", meta_lbl), Paragraph("Software (AI/ML + Geospatial GIS + Offline PWA + Multilingual)", meta_val)],
                [Paragraph("<b>• Team ID:</b>", meta_lbl), Paragraph("[Your Team ID]", meta_val)],
                [Paragraph("<b>• Team Name:</b>", meta_lbl), Paragraph("[Your Team Name - Registered on Portal]", meta_val)],
            ], colWidths=[160, 520])
        ]
    ]
    t_table = Table(t_box_data, colWidths=[720])
    t_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.white),
        ('BOX', (0,0), (-1,-1), 1.5, c_teal),
        ('PADDING', (0,0), (-1,-1), 8),
        ('ALIGN', (0,0), (-1,2), 'CENTER'),
    ]))
    story.append(t_table)
    story.append(PageBreak())

    # -------------------------------------------------------------
    # SLIDE 2: PROPOSED SOLUTION
    # -------------------------------------------------------------
    story.append(make_slide_header("Idea Title: NER-LRP Proposed Solution", 2))
    story.append(Spacer(1, 8))

    col1 = [
        Paragraph("<b>Detailed Explanation of Proposed Solution</b>", card_h_style),
        Paragraph("• <b>Sense-Predict-Decide-Act Control Tower:</b> Integrated regional command platform monitoring 4,820 km of critical mountain highways across all 8 Northeast states.", body_style),
        Paragraph("• <b>Real Live Data Ingestion:</b> Automatically ingests live precipitation & 48h rain from Open-Meteo/IMD APIs; integrates real NASA SRTM 30m slope gradients (e.g. 38.5° on NH-10).", body_style),
        Paragraph("• <b>AI Hazard Risk Engine:</b> Gradient Boosting ML model predicts 0-100 hazard scores with explainable factor contributions (Rain vs Slope vs Soil vs History).", body_style),
        Paragraph("• <b>Dynamic Route Re-planning:</b> Recalculates safe bypasses (e.g. NH-27 Haflong bypass around blocked NH-6 Sonapur) with exact ETA deltas.", body_style)
    ]

    col2 = [
        Paragraph("<b>How It Addresses Problem & Key Innovations</b>", card_h_teal),
        Paragraph("• <b>Cargo Criticality Weighting (Novel):</b> 85% safety priority for Emergency Medicines & Oxygen vs 25% for Commercial Bulk. Eliminates life-or-death stranding.", body_style),
        Paragraph("• <b>Offline-First Field Officer PWA:</b> Allows patrol officers in 0-connectivity mountain valleys to submit geo-photo reports cached in IndexedDB with auto-sync upon reconnect.", body_style),
        Paragraph("• <b>Multilingual Emergency Broadcaster:</b> Translates road warnings into 6 regional languages (Assamese, Mizo, Manipuri, Bengali, Hindi, English) with audible voice synthesis.", body_style),
        Paragraph("• <b>1-Click Disruption Simulator:</b> Enables live evaluators to trigger instant landslides (NH-6, NH-10, NH-29) and observe real-time rerouting in action.", body_style)
    ]

    s2_table = Table([[col1, col2]], colWidths=[355, 355])
    s2_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), c_light),
        ('BOX', (0,0), (-1,-1), 0.5, c_border),
        ('INNERGRID', (0,0), (-1,-1), 0.5, c_border),
        ('PADDING', (0,0), (-1,-1), 8),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(s2_table)
    story.append(PageBreak())

    # -------------------------------------------------------------
    # SLIDE 3: TECHNICAL APPROACH
    # -------------------------------------------------------------
    story.append(make_slide_header("Technical Approach & Architecture Workflow", 3))
    story.append(Spacer(1, 8))

    t_box1 = [
        Paragraph("<b>Backend & GIS Pipeline</b>", card_h_style),
        Paragraph("• Python FastAPI (Async REST API)<br/>• Open-Meteo & IMD Live Weather Feeds<br/>• SRTM DEM 30m Slope Gradient Data<br/>• SQLite ACID DB with Spatial Schemas", body_style)
    ]
    t_box2 = [
        Paragraph("<b>AI / ML & Routing Engine</b>", card_h_teal),
        Paragraph("• Scikit-Learn Gradient Boosting ML<br/>• Multi-variable Hazard Score (0-100)<br/>• Dijkstra Multi-Objective Graph Router<br/>• Cargo Priority Cost Functions", body_style)
    ]
    t_box3 = [
        Paragraph("<b>Frontend & Offline PWA</b>", ParagraphStyle('CE', parent=card_h_style, textColor=colors.HexColor("#059669"))),
        Paragraph("• React 18 + Vite + Tailwind Dark UI<br/>• Leaflet GIS + Dynamic Map Layers<br/>• IndexedDB Offline Local Storage Queue<br/>• Web Speech Synthesis API (6 Languages)", body_style)
    ]

    s3_top = Table([[t_box1, t_box2, t_box3]], colWidths=[236, 236, 236])
    s3_top.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), c_light),
        ('BOX', (0,0), (-1,-1), 0.5, c_border),
        ('INNERGRID', (0,0), (-1,-1), 0.5, c_border),
        ('PADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(s3_top)
    story.append(Spacer(1, 6))

    wf_content = [
        Paragraph("<b>End-to-End Implementation Methodology Workflow:</b>", ParagraphStyle('W', parent=body_style, fontName='Helvetica-Bold', textColor=c_navy)),
        Paragraph("<b>1. Live Ingestion:</b> Open-Meteo queries rain/wind/soil moisture across 15+ NER stations + SRTM slope profiles ➔ <b>2. ML Hazard Scoring:</b> GBDT model predicts 0-100 risk with explainable factors ➔ <b>3. Cargo-Aware Routing:</b> Diverts medical shipments around risky segments with delta ETA ➔ <b>4. Control Tower & Field PWA:</b> Visualizes live GIS map & syncs offline damage reports ➔ <b>5. Multilingual Dispatch:</b> Broadcasts voice/text alerts across 6 regional languages.", body_style)
    ]
    s3_bot = Table([[wf_content]], colWidths=[720])
    s3_bot.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.white),
        ('BOX', (0,0), (-1,-1), 0.5, c_teal),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(s3_bot)
    story.append(PageBreak())

    # -------------------------------------------------------------
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # -------------------------------------------------------------
    story.append(make_slide_header("Feasibility, Viability & Risk Mitigation", 4))
    story.append(Spacer(1, 8))

    f1 = [
        Paragraph("<b>Feasibility & Validation</b>", card_h_style),
        Paragraph("• <b>Open Data Architecture:</b> Built entirely on accessible open data (Open-Meteo, OpenStreetMap, NASA SRTM). Zero licensing cost.", body_style),
        Paragraph("• <b>Working Prototype:</b> Fully functional full-stack app verified with 21 passing automated unit tests.", body_style),
        Paragraph("• <b>Corridor-Scale Focus:</b> Validated on critical lifelines: NH-10 (Sikkim), NH-6 (Meghalaya), NH-29 (Nagaland), NH-27 (Assam).", body_style)
    ]
    f2 = [
        Paragraph("<b>Potential Risks & Challenges</b>", ParagraphStyle('R', parent=card_h_style, textColor=colors.HexColor("#ea580c"))),
        Paragraph("• <b>Govt Fleet Feeds Locked:</b> Real-time ULIP telematics require ministry authorizations.", body_style),
        Paragraph("• <b>Cellular Blackouts:</b> Mountain ghats frequently suffer zero 4G/5G mobile coverage.", body_style),
        Paragraph("• <b>Extreme Himalayan Terrain:</b> Sudden cloudbursts cause rapid slope failures.", body_style)
    ]
    f3 = [
        Paragraph("<b>Mitigation Strategies</b>", card_h_teal),
        Paragraph("• <b>Modular ULIP Adapter:</b> Standard GeoJSON interface ready to plug into ULIP when NDA is approved.", body_style),
        Paragraph("• <b>Offline-First PWA:</b> Caches reports in IndexedDB; auto-syncs idempotently upon signal recovery.", body_style),
        Paragraph("• <b>Hybrid Physics + ML:</b> Combines empirical IRC slope thresholds with GBDT machine learning.", body_style)
    ]

    s4_table = Table([[f1, f2, f3]], colWidths=[236, 236, 236])
    s4_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), c_light),
        ('BOX', (0,0), (-1,-1), 0.5, c_border),
        ('INNERGRID', (0,0), (-1,-1), 0.5, c_border),
        ('PADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(s4_table)
    story.append(PageBreak())

    # -------------------------------------------------------------
    # SLIDE 5: IMPACT AND BENEFITS
    # -------------------------------------------------------------
    story.append(make_slide_header("Impact, Benefits & Stakeholder Value", 5))
    story.append(Spacer(1, 8))

    q1 = [
        Paragraph("<b>Target Audience Impact</b>", card_h_style),
        Paragraph("• <b>MDoNER & SDMAs:</b> Unified regional command tower across all 8 Northeast states for coordinated emergency planning.<br/>• <b>Hospitals:</b> Prevents stock-out of emergency oxygen, anti-venom, and vaccines.<br/>• <b>Drivers:</b> Hands-free voice alerts prevent 48h mountain jam entrapment.", body_style)
    ]
    q2 = [
        Paragraph("<b>Economic Benefits</b>", card_h_teal),
        Paragraph("• <b>Inflation Suppression:</b> Prevents local price spikes (>20%) in isolated districts.<br/>• <b>Fuel & Delay Savings:</b> Cuts vehicle turnaround times by 15-25% via planned detours.<br/>• <b>Asset Protection:</b> Protects multi-crore truck assets from sudden rockfall crushes.", body_style)
    ]
    q3 = [
        Paragraph("<b>Social & Environmental Benefits</b>", card_h_teal),
        Paragraph("• <b>Supply Reliability:</b> Uninterrupted food grains & PDS rations to remote tribal hamlets.<br/>• <b>Emission Reduction:</b> Eliminates thousands of hours of diesel truck idling during blockades.<br/>• <b>Grassroots Reporting:</b> Empowers local scouts to report hazards directly.", body_style)
    ]
    q4 = [
        Paragraph("<b>System Resilience & PM GatiShakti</b>", card_h_style),
        Paragraph("• <b>Emergency Corridors:</b> Uncovers secondary state bypasses during national highway closures.<br/>• <b>GatiShakti Alignment:</b> Assists infrastructure planners in reinforcing high-risk slopes.<br/>• <b>Himalayan Scalability:</b> Replicable across Uttarakhand, HP, and J&K.", body_style)
    ]

    s5_table = Table([[q1, q2], [q3, q4]], colWidths=[355, 355])
    s5_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), c_light),
        ('BOX', (0,0), (-1,-1), 0.5, c_border),
        ('INNERGRID', (0,0), (-1,-1), 0.5, c_border),
        ('PADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(s5_table)
    story.append(PageBreak())

    # -------------------------------------------------------------
    # SLIDE 6: RESEARCH AND REFERENCES
    # -------------------------------------------------------------
    story.append(make_slide_header("Research, Citations & Policy Frameworks", 6))
    story.append(Spacer(1, 8))

    r_col1 = [
        Paragraph("<b>Geospatial Data & Scientific Literature</b>", card_h_style),
        Paragraph("• <b>Geological Survey of India (GSI):</b> National Landslide Susceptibility Mapping (NLSM) and Landslide Hazard Zonation (LHZ) protocols for Northeast India.<br/>"
                  "• <b>Academic ML Studies (Eastern Himalayas):</b> Gradient Boosting (XGBoost/LightGBM) models for landslide susceptibility in NE India (Dibang Valley & Nagaland, achieving AUC ~0.96 and ~0.89).<br/>"
                  "• <b>NASA SRTM 30m DEM:</b> High-resolution digital elevation data used for terrain slope and drainage gradient modeling across Himalayan highways.<br/>"
                  "• <b>Open-Meteo & IMD APIs:</b> Public meteorological endpoints providing live precipitation, 48h cumulative rainfall, and Doppler weather radar data.", body_style)
    ]

    r_col2 = [
        Paragraph("<b>Government Policies & Frameworks</b>", card_h_teal),
        Paragraph("• <b>PM GatiShakti National Master Plan:</b> Ministry of Commerce & Industry digital planning mandate for multimodal connectivity and logistics infrastructure resilience.<br/>"
                  "• <b>High-Level Task Force (HLTF) NER Report:</b> Official MDoNER reports detailing transport bottlenecks, logistics costs, and the 20% price inflation in remote hill states.<br/>"
                  "• <b>ULIP (Unified Logistics Interface Platform):</b> National Logistics Policy framework integrating multi-ministry vehicle telematics and digital transport documentation.<br/>"
                  "• <b>OpenStreetMap & OSRM Engine:</b> Open road graph datasets customized for heavy truck routing and slope penalty costs.", body_style)
    ]

    s6_table = Table([[r_col1, r_col2]], colWidths=[355, 355])
    s6_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), c_light),
        ('BOX', (0,0), (-1,-1), 0.5, c_border),
        ('INNERGRID', (0,0), (-1,-1), 0.5, c_border),
        ('PADDING', (0,0), (-1,-1), 8),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(s6_table)

    doc.build(story)
    print(f"[SUCCESS] PDF presentation successfully generated at: {filename}")


if __name__ == "__main__":
    create_pptx()
    create_pdf_presentation()
